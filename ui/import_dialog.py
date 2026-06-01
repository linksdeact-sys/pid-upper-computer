#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
数据导入对话框
用于导入CSV和JSON格式的数据
"""

import os
from PyQt5.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QGroupBox,
    QLabel, QPushButton, QLineEdit, QFileDialog,
    QComboBox, QCheckBox, QDialogButtonBox,
    QMessageBox, QTableWidget, QTableWidgetItem,
    QHeaderView
)
from PyQt5.QtCore import Qt


class ImportDialog(QDialog):
    """数据导入对话框"""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("导入数据")
        self.setMinimumSize(600, 500)
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)

        # 导入的数据
        self.imported_data = []

        # 初始化界面
        self._init_ui()

    def _init_ui(self):
        """初始化界面"""
        layout = QVBoxLayout(self)
        layout.setSpacing(16)

        # 文件选择组
        file_group = QGroupBox("文件选择")
        file_layout = QHBoxLayout()

        self.file_path_edit = QLineEdit()
        self.file_path_edit.setPlaceholderText("选择数据文件...")
        self.file_path_edit.setReadOnly(True)

        browse_btn = QPushButton("浏览...")
        browse_btn.clicked.connect(self._browse_file)

        file_layout.addWidget(self.file_path_edit)
        file_layout.addWidget(browse_btn)
        file_group.setLayout(file_layout)
        layout.addWidget(file_group)

        # 导入选项组
        options_group = QGroupBox("导入选项")
        options_layout = QVBoxLayout()

        # 文件格式
        format_layout = QHBoxLayout()
        format_layout.addWidget(QLabel("文件格式:"))
        self.format_combo = QComboBox()
        self.format_combo.addItems(["自动检测", "CSV", "JSON", "TXT", "Excel"])
        format_layout.addWidget(self.format_combo)
        format_layout.addStretch()
        options_layout.addLayout(format_layout)

        # 编码选择
        encoding_layout = QHBoxLayout()
        encoding_layout.addWidget(QLabel("文件编码:"))
        self.encoding_combo = QComboBox()
        self.encoding_combo.addItems(["UTF-8", "GBK", "GB2312", "ISO-8859-1"])
        encoding_layout.addWidget(self.encoding_combo)
        encoding_layout.addStretch()
        options_layout.addLayout(encoding_layout)

        # 导入选项
        self.clear_existing_check = QCheckBox("导入前清空现有数据")
        self.clear_existing_check.setChecked(True)
        options_layout.addWidget(self.clear_existing_check)

        options_group.setLayout(options_layout)
        layout.addWidget(options_group)

        # 预览表格
        preview_group = QGroupBox("数据预览")
        preview_layout = QVBoxLayout()

        self.preview_table = QTableWidget()
        self.preview_table.setAlternatingRowColors(True)
        self.preview_table.setSelectionBehavior(QTableWidget.SelectRows)
        preview_layout.addWidget(self.preview_table)

        # 预览信息
        self.preview_info_label = QLabel("请选择数据文件")
        self.preview_info_label.setStyleSheet("color: #8E8E93;")
        preview_layout.addWidget(self.preview_info_label)

        preview_group.setLayout(preview_layout)
        layout.addWidget(preview_group)

        # 按钮
        button_box = QDialogButtonBox(
            QDialogButtonBox.Ok | QDialogButtonBox.Cancel
        )
        button_box.button(QDialogButtonBox.Ok).setText("导入")
        button_box.button(QDialogButtonBox.Cancel).setText("取消")
        button_box.accepted.connect(self._on_import)
        button_box.rejected.connect(self.reject)
        layout.addWidget(button_box)

    def _browse_file(self):
        """浏览文件"""
        filepath, _ = QFileDialog.getOpenFileName(
            self, "选择数据文件", "",
            "CSV文件 (*.csv);;JSON文件 (*.json);;文本文件 (*.txt);;Excel文件 (*.xlsx);;所有文件 (*)"
        )

        if filepath:
            self.file_path_edit.setText(filepath)
            self._preview_file(filepath)

    def _preview_file(self, filepath):
        """预览文件内容"""
        try:
            # 检测文件格式
            if filepath.endswith('.csv'):
                format_type = 'CSV'
            elif filepath.endswith('.json'):
                format_type = 'JSON'
            elif filepath.endswith('.txt'):
                format_type = 'TXT'
            elif filepath.endswith('.xlsx'):
                format_type = 'Excel'
            else:
                format_type = '未知'

            # 更新格式选择
            if format_type == 'CSV':
                self.format_combo.setCurrentText("CSV")
            elif format_type == 'JSON':
                self.format_combo.setCurrentText("JSON")
            elif format_type == 'TXT':
                self.format_combo.setCurrentText("TXT")
            elif format_type == 'Excel':
                self.format_combo.setCurrentText("Excel")

            # 读取文件内容
            encoding = self.encoding_combo.currentText()

            if format_type == 'CSV':
                self._preview_csv(filepath, encoding)
            elif format_type == 'JSON':
                self._preview_json(filepath, encoding)
            elif format_type == 'TXT':
                self._preview_txt(filepath, encoding)
            elif format_type == 'Excel':
                self._preview_excel(filepath)
            else:
                self.preview_info_label.setText("不支持的文件格式")

        except Exception as e:
            self.preview_info_label.setText(f"预览失败: {str(e)}")
            QMessageBox.critical(self, "错误", f"读取文件失败: {str(e)}")

    def _preview_csv(self, filepath, encoding):
        """预览CSV文件"""
        import csv

        with open(filepath, 'r', encoding=encoding) as f:
            reader = csv.reader(f)
            rows = list(reader)

            if not rows:
                self.preview_info_label.setText("文件为空")
                return

            # 设置表头
            headers = rows[0]
            self.preview_table.setColumnCount(len(headers))
            self.preview_table.setHorizontalHeaderLabels(headers)

            # 设置数据行数（最多显示100行）
            data_rows = rows[1:101]
            self.preview_table.setRowCount(len(data_rows))

            # 填充数据
            for i, row in enumerate(data_rows):
                for j, value in enumerate(row):
                    if j < len(headers):
                        self.preview_table.setItem(i, j, QTableWidgetItem(value))

            # 更新信息
            total_rows = len(rows) - 1
            self.preview_info_label.setText(
                f"共 {total_rows} 行数据，显示前 {len(data_rows)} 行"
            )

    def _preview_json(self, filepath, encoding):
        """预览JSON文件"""
        import json

        with open(filepath, 'r', encoding=encoding) as f:
            data = json.load(f)

            # 检查数据格式
            if isinstance(data, dict):
                if 'data' in data:
                    items = data['data']
                else:
                    items = [data]
            elif isinstance(data, list):
                items = data
            else:
                self.preview_info_label.setText("不支持的JSON格式")
                return

            if not items:
                self.preview_info_label.setText("文件为空")
                return

            # 获取所有键作为表头
            headers = list(items[0].keys())
            self.preview_table.setColumnCount(len(headers))
            self.preview_table.setHorizontalHeaderLabels(headers)

            # 设置数据行数（最多显示100行）
            display_items = items[:100]
            self.preview_table.setRowCount(len(display_items))

            # 填充数据
            for i, item in enumerate(display_items):
                for j, header in enumerate(headers):
                    value = str(item.get(header, ''))
                    self.preview_table.setItem(i, j, QTableWidgetItem(value))

            # 更新信息
            self.preview_info_label.setText(
                f"共 {len(items)} 条数据，显示前 {len(display_items)} 条"
            )

    def _preview_txt(self, filepath, encoding):
        """预览TXT文件"""
        with open(filepath, 'r', encoding=encoding) as f:
            lines = f.readlines()

            if not lines:
                self.preview_info_label.setText("文件为空")
                return

            # 解析表头
            headers = lines[0].strip().split('\t')
            self.preview_table.setColumnCount(len(headers))
            self.preview_table.setHorizontalHeaderLabels(headers)

            # 设置数据行数（最多显示100行）
            data_lines = lines[1:101]
            self.preview_table.setRowCount(len(data_lines))

            # 填充数据
            for i, line in enumerate(data_lines):
                values = line.strip().split('\t')
                for j, value in enumerate(values):
                    if j < len(headers):
                        self.preview_table.setItem(i, j, QTableWidgetItem(value))

            # 更新信息
            total_rows = len(lines) - 1
            self.preview_info_label.setText(
                f"共 {total_rows} 行数据，显示前 {len(data_lines)} 行"
            )

    def _preview_excel(self, filepath):
        """预览Excel文件"""
        try:
            import openpyxl

            wb = openpyxl.load_workbook(filepath, read_only=True)
            ws = wb.active

            # 获取表头
            headers = []
            for cell in ws[1]:
                headers.append(str(cell.value) if cell.value else '')

            self.preview_table.setColumnCount(len(headers))
            self.preview_table.setHorizontalHeaderLabels(headers)

            # 获取数据行数
            total_rows = ws.max_row - 1
            display_rows = min(100, total_rows)

            self.preview_table.setRowCount(display_rows)

            # 填充数据
            for i, row in enumerate(ws.iter_rows(min_row=2, max_row=display_rows + 1)):
                for j, cell in enumerate(row):
                    if j < len(headers):
                        value = str(cell.value) if cell.value is not None else ''
                        self.preview_table.setItem(i, j, QTableWidgetItem(value))

            wb.close()

            # 更新信息
            self.preview_info_label.setText(
                f"共 {total_rows} 行数据，显示前 {display_rows} 行"
            )

        except ImportError:
            self.preview_info_label.setText("预览Excel需要安装openpyxl库")
        except Exception as e:
            self.preview_info_label.setText(f"预览失败: {str(e)}")

    def _on_import(self):
        """导入数据"""
        filepath = self.file_path_edit.text()
        if not filepath:
            QMessageBox.warning(self, "警告", "请先选择数据文件")
            return

        try:
            # 获取导入选项
            format_type = self.format_combo.currentText()
            encoding = self.encoding_combo.currentText()
            clear_existing = self.clear_existing_check.isChecked()

            # 检测格式
            if format_type == "自动检测":
                if filepath.endswith('.csv'):
                    format_type = "CSV"
                elif filepath.endswith('.json'):
                    format_type = "JSON"
                elif filepath.endswith('.txt'):
                    format_type = "TXT"
                elif filepath.endswith('.xlsx'):
                    format_type = "Excel"
                else:
                    QMessageBox.warning(self, "警告", "无法自动检测文件格式")
                    return

            # 读取数据
            if format_type == "CSV":
                self._import_csv(filepath, encoding)
            elif format_type == "JSON":
                self._import_json(filepath, encoding)
            elif format_type == "TXT":
                self._import_txt(filepath, encoding)
            elif format_type == "Excel":
                self._import_excel(filepath)

            if self.imported_data:
                self.accept()
            else:
                QMessageBox.warning(self, "警告", "没有导入任何数据")

        except Exception as e:
            QMessageBox.critical(self, "错误", f"导入失败: {str(e)}")

    def _import_csv(self, filepath, encoding):
        """导入CSV数据"""
        import csv

        self.imported_data = []
        with open(filepath, 'r', encoding=encoding) as f:
            reader = csv.DictReader(f)
            for row in reader:
                self.imported_data.append(row)

    def _import_json(self, filepath, encoding):
        """导入JSON数据"""
        import json

        with open(filepath, 'r', encoding=encoding) as f:
            data = json.load(f)

            if isinstance(data, dict):
                if 'data' in data:
                    self.imported_data = data['data']
                else:
                    self.imported_data = [data]
            elif isinstance(data, list):
                self.imported_data = data
            else:
                self.imported_data = []

    def _import_txt(self, filepath, encoding):
        """导入TXT数据"""
        self.imported_data = []
        with open(filepath, 'r', encoding=encoding) as f:
            lines = f.readlines()

            if not lines:
                return

            # 解析表头
            headers = lines[0].strip().split('\t')

            # 解析数据
            for line in lines[1:]:
                values = line.strip().split('\t')
                if len(values) >= len(headers):
                    row = {}
                    for i, header in enumerate(headers):
                        row[header] = values[i]
                    self.imported_data.append(row)

    def _import_excel(self, filepath):
        """导入Excel数据"""
        try:
            import openpyxl

            wb = openpyxl.load_workbook(filepath, read_only=True)
            ws = wb.active

            # 获取表头
            headers = []
            for cell in ws[1]:
                headers.append(str(cell.value) if cell.value else '')

            # 获取数据
            self.imported_data = []
            for row in ws.iter_rows(min_row=2):
                row_data = {}
                for i, cell in enumerate(row):
                    if i < len(headers):
                        row_data[headers[i]] = str(cell.value) if cell.value is not None else ''
                self.imported_data.append(row_data)

            wb.close()

        except ImportError:
            QMessageBox.warning(self, "警告", "导入Excel需要安装openpyxl库\n请运行: pip install openpyxl")
        except Exception as e:
            QMessageBox.critical(self, "错误", f"导入Excel失败: {str(e)}")

    def get_imported_data(self):
        """获取导入的数据"""
        return self.imported_data

    def should_clear_existing(self):
        """是否清空现有数据"""
        return self.clear_existing_check.isChecked()
